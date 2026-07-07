from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor


REPO_ROOT = Path(__file__).resolve().parents[1]
SKELETON_ROOT = REPO_ROOT / "templates" / "skeletons" / "Lenovo.XCC"
OUTPUT_PATH = SKELETON_ROOT / "XCC-SDT-Collector.docx"
END_DOCUMENT_OUTPUT_PATH = SKELETON_ROOT / "XCC-EndDocument.docx"


SECTIONS = [
    ("Collection Status", ["LNV.Lenovo.XCC.Target[TargetName].Tables.CollectionStatus"]),
    ("System Overview", ["LNV.Lenovo.XCC.Target[TargetName].Summary"]),
    (
        "Hardware Inventory",
        [
            "LNV.Lenovo.XCC.Target[TargetName].Tables.Managers",
            "LNV.Lenovo.XCC.Target[TargetName].Tables.Chassis",
            "LNV.Lenovo.XCC.Target[TargetName].Tables.Firmware",
            "LNV.Lenovo.XCC.Target[TargetName].Tables.Processors",
            "LNV.Lenovo.XCC.Target[TargetName].Tables.Memory",
        ],
    ),
    (
        "Access and Security",
        [
            "LNV.Lenovo.XCC.Target[TargetName].Tables.Security",
            "LNV.Lenovo.XCC.Target[TargetName].Tables.Users",
            "LNV.Lenovo.XCC.Target[TargetName].Tables.Bios",
            "LNV.Lenovo.XCC.Target[TargetName].Tables.SecureBoot",
        ],
    ),
    (
        "Management Interfaces",
        [
            "LNV.Lenovo.XCC.Target[TargetName].Tables.ManagerNetworkProtocol",
            "LNV.Lenovo.XCC.Target[TargetName].Tables.ManagerEthernetInterfaces",
            "LNV.Lenovo.XCC.Target[TargetName].Tables.ManagerHostInterfaces",
            "LNV.Lenovo.XCC.Target[TargetName].Tables.ManagerSerialInterfaces",
            "LNV.Lenovo.XCC.Target[TargetName].Tables.VirtualMedia",
            "LNV.Lenovo.XCC.Target[TargetName].Tables.ManagementEthernetPorts",
        ],
    ),
    (
        "Server Connectivity",
        [
            "LNV.Lenovo.XCC.Target[TargetName].Tables.NetworkInterfaces",
            "LNV.Lenovo.XCC.Target[TargetName].Tables.HostEthernetPorts",
            "LNV.Lenovo.XCC.Target[TargetName].Tables.HostFcPorts",
            "LNV.Lenovo.XCC.Target[TargetName].Tables.ServerPortInventory",
        ],
    ),
    (
        "SE350 Embedded Switch-Board Evidence",
        [
            "LNV.Lenovo.XCC.Target[TargetName].Tables.SystemNetworkInterfaces",
            "LNV.Lenovo.XCC.Target[TargetName].Tables.NetworkAdapterPorts",
            "LNV.Lenovo.XCC.Target[TargetName].Tables.NetworkAdapterNetworkPorts",
            "LNV.Lenovo.XCC.Target[TargetName].Tables.NetworkDeviceFunctions",
        ],
    ),
    (
        "PCIe and Chassis",
        [
            "LNV.Lenovo.XCC.Target[TargetName].Tables.PcieDevices",
            "LNV.Lenovo.XCC.Target[TargetName].Tables.PcieFunctions",
            "LNV.Lenovo.XCC.Target[TargetName].Tables.ChassisNetworkAdapters",
            "LNV.Lenovo.XCC.Target[TargetName].Tables.ChassisPcieSlots",
        ],
    ),
    (
        "Storage and Environment",
        [
            "LNV.Lenovo.XCC.Target[TargetName].Tables.StorageControllers",
            "LNV.Lenovo.XCC.Target[TargetName].Tables.StorageDrives",
            "LNV.Lenovo.XCC.Target[TargetName].Tables.StorageVolumes",
            "LNV.Lenovo.XCC.Target[TargetName].Tables.Power",
            "LNV.Lenovo.XCC.Target[TargetName].Tables.Thermal",
            "LNV.Lenovo.XCC.Target[TargetName].Tables.ChassisSensors",
            "LNV.Lenovo.XCC.Target[TargetName].Tables.ChassisEnvironment",
            "LNV.Lenovo.XCC.Target[TargetName].Tables.ChassisPowerSubsystem",
            "LNV.Lenovo.XCC.Target[TargetName].Tables.ChassisThermalSubsystem",
        ],
    ),
    (
        "Topology Graph Inputs",
        [
            "LNV.Lenovo.XCC.Target[TargetName].Diagrams.ServerPortGraph",
            "LNV.Lenovo.XCC.Target[TargetName].Diagrams.StorageConnectivityGraph",
        ],
    ),
    (
        "Appendix: Redfish Port Evidence",
        [
            "LNV.Lenovo.XCC.Target[TargetName].Appendix.HostEthernetPortEvidence",
            "LNV.Lenovo.XCC.Target[TargetName].Appendix.ServerPortInventoryEvidence",
            "LNV.Lenovo.XCC.Target[TargetName].Appendix.SystemNetworkInterfaceEvidence",
            "LNV.Lenovo.XCC.Target[TargetName].Appendix.NetworkDeviceFunctionEvidence",
        ],
    ),
    ("Event Log", ["LNV.Lenovo.XCC.Target[TargetName].Tables.EventLog"]),
]


END_DOCUMENT_SECTIONS = [
    ("Vital Product Data", ["LNV.Lenovo.XCC.Target[TargetName].Document.VitalProductDataGrouped"]),
    ("Firmware Summary", ["LNV.Lenovo.XCC.Target[TargetName].Document.FirmwareSummary"]),
    (
        "Hardware Configuration",
        [
            "LNV.Lenovo.XCC.Target[TargetName].Document.HardwareSummary",
            "LNV.Lenovo.XCC.Target[TargetName].Document.MemoryExceptions",
        ],
    ),
    (
        "Network Configuration",
        [
            "LNV.Lenovo.XCC.Target[TargetName].Document.ManagementEthernetPorts",
            "LNV.Lenovo.XCC.Target[TargetName].Document.HostEthernetPorts",
            "LNV.Lenovo.XCC.Target[TargetName].Document.HostFcPorts",
        ],
    ),
]


TAG_LABELS = {
    "LNV.Lenovo.XCC.Target[TargetName].Document.VitalProductDataGrouped": "Vital Product Data",
    "LNV.Lenovo.XCC.Target[TargetName].Document.FirmwareSummary": "Firmware Summary",
    "LNV.Lenovo.XCC.Target[TargetName].Document.HardwareSummary": "Hardware Summary",
    "LNV.Lenovo.XCC.Target[TargetName].Document.MemoryExceptions": "Memory Exceptions",
    "LNV.Lenovo.XCC.Target[TargetName].Document.ManagementEthernetPorts": "XCC Interfaces",
    "LNV.Lenovo.XCC.Target[TargetName].Document.HostEthernetPorts": "Host Ethernet Ports",
    "LNV.Lenovo.XCC.Target[TargetName].Document.HostFcPorts": "Host Fibre Channel Ports",
}


def add_token_paragraph(document, tag):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(f"<<SDT:{tag}>>")
    run.font.name = "Arial"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string("475467")
    document.add_paragraph()


def configure_document(document):
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9)

    for style_name, size, color in [
        ("Title", 20, "1F2937"),
        ("Heading 1", 14, "1F2937"),
        ("Heading 2", 11, "344054"),
    ]:
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True


def build_collector_blueprint():
    document = Document()
    configure_document(document)

    title = document.add_paragraph()
    title.style = document.styles["Title"]
    title.add_run("Lenovo XCC Collector Blueprint")

    subtitle = document.add_paragraph()
    subtitle.add_run("Contract-driven skeleton for Lenovo XCC inventory, management, and server-port topology projections.")

    for section_index, (heading, tags) in enumerate(SECTIONS):
        if section_index in {5, 8}:
            document.add_page_break()
        document.add_heading(heading, level=1)
        for tag in tags:
            label = tag.split(".")[-1]
            document.add_heading(label, level=2)
            add_token_paragraph(document, tag)

    for section in document.sections:
        footer = section.footer.paragraphs[0]
        footer.text = "Lenovo XCC collector skeleton - generated from assembler contract SDT mappings"
        footer.runs[0].font.name = "Arial"
        footer.runs[0].font.size = Pt(8)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


def build_end_document():
    document = Document()
    configure_document(document)

    title = document.add_paragraph()
    title.style = document.styles["Title"]
    title.add_run("Lenovo XCC As-Built Summary")

    subtitle = document.add_paragraph()
    subtitle.add_run("Customer-facing hardware, firmware, management access, and server port projections.")

    for section_index, (heading, tags) in enumerate(END_DOCUMENT_SECTIONS):
        if section_index == 3:
            document.add_page_break()
        document.add_heading(heading, level=1)
        for tag in tags:
            label = TAG_LABELS.get(tag, tag.split(".")[-1])
            document.add_heading(label, level=2)
            add_token_paragraph(document, tag)

    for section in document.sections:
        footer = section.footer.paragraphs[0]
        footer.text = "Lenovo XCC as-built summary - generated from assembler contract SDT mappings"
        footer.runs[0].font.name = "Arial"
        footer.runs[0].font.size = Pt(8)

    END_DOCUMENT_OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(END_DOCUMENT_OUTPUT_PATH)
    print(END_DOCUMENT_OUTPUT_PATH)


def main():
    build_collector_blueprint()
    build_end_document()


if __name__ == "__main__":
    main()
